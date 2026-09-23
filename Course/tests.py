import tempfile
import importlib
import json
import requests
from datetime import timedelta
from io import BytesIO, StringIO
from unittest.mock import Mock, patch

from django.contrib import admin
from django.apps import apps
from django.contrib.auth.models import Permission, User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.db import IntegrityError, transaction
from django.test import RequestFactory, TestCase
from django.test.utils import override_settings
from django.urls import reverse
from django.utils import timezone

from .models import (
    AdditionalLessons,
    DirectionStudy,
    LearnGroup,
    LessonSolution,
    LessonSolutionSubmission,
    Schedule,
    Student,
)
from .admin import (
    ActiveGroupListFilter,
    ActiveStudentListFilter,
    AttentionNeededListFilter,
    LessonSolutionAdmin,
    MyGroupsListFilter,
    RecentSolutionListFilter,
    ScheduleAdmin,
    WithoutCommentListFilter,
)
from .views import get_accessible_schedule_ids


class TimetableTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(
            "schedule-user",
            email="student@example.com",
            password="password",
        )
        self.direction = DirectionStudy.objects.create(title="Python")
        self.student = Student.objects.create(
            user=self.user,
            contact="@student",
            is_learned=True,
            groups_id=100,
        )
        self.group = LearnGroup.objects.create(
            pk=100,
            title="Учебная группа",
            is_studies=True,
            teacher=self.student,
        )
        self.student.direction.add(self.direction)
        AdditionalLessons.objects.create(group=self.group, amount=10)

        for number in range(25):
            Schedule.objects.create(
                theme=f"Урок {number}",
                plan="План",
                lesson_materials=f"Материалы {number}",
                lesson_type="Новая тема" if number % 2 else "Практика",
                direction=self.direction,
            )

    def test_account_link_uses_one_to_one_relation(self):
        duplicate_user = User.objects.create_user(
            'duplicate-account',
            email='duplicate@example.com',
            password='password',
        )
        duplicate_profile = Student.objects.create(
            user=duplicate_user,
            contact='@duplicate',
            groups=self.group,
        )

        self.assertEqual(Student.objects.for_user(self.user), self.student)
        self.assertEqual(
            Student.objects.for_user(duplicate_user),
            duplicate_profile,
        )

    def test_username_lookup_uses_canonical_user_username(self):
        canonical_user = User.objects.create_user(
            'canonical-account',
            email='canonical@example.com',
            password='password',
        )
        canonical_profile = Student.objects.create(
            user=canonical_user,
            contact='@canonical',
            groups=self.group,
        )

        self.assertEqual(
            Student.objects.for_username(canonical_user.username),
            canonical_profile,
        )
        self.assertIsNone(Student.objects.for_username('old-display-name'))

    def test_profile_cannot_be_created_without_account(self):
        with self.assertRaises(IntegrityError), transaction.atomic():
            Student.objects.create(
                contact='@missing-account',
                groups=self.group,
            )

    def test_link_audit_never_prints_passwords(self):
        self.user.set_password('private-account-password')
        self.user.save(update_fields=('password',))
        output = StringIO()

        call_command('audit_user_student_links', '--details', stdout=output)

        report = output.getvalue()
        self.assertIn('Связанных профилей', report)
        self.assertNotIn('private-account-password', report)

    def test_registration_creates_linked_account_and_profile_atomically(self):
        LearnGroup.objects.create(
            pk=2,
            title='Группа новых заявок',
            is_studies=False,
            teacher=self.student,
        )

        response = self.client.post(
            reverse('home'),
            {
                'name': 'new-linked-student',
                'contact': '@new-linked-student',
                'email': 'new-linked@example.com',
                'password': 'registration-password',
            },
            REMOTE_ADDR='192.0.2.10',
        )

        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json()['success'])
        user = User.objects.get(username='new-linked-student')
        student = Student.objects.get(user=user)
        self.assertEqual(student.user.username, user.username)
        self.assertEqual(student.user.email, 'new-linked@example.com')
        self.assertTrue(user.check_password('registration-password'))

    def test_timetable_requires_learned_student(self):
        anonymous_response = self.client.get(reverse("schedule"))
        another_user = User.objects.create_user(
            "not-a-student",
            password="password",
        )
        self.client.force_login(another_user)
        unrelated_response = self.client.get(reverse("schedule"))

        self.assertEqual(anonymous_response.status_code, 302)
        self.assertRedirects(unrelated_response, reverse("home"))

    def test_timetable_uses_server_side_pagination(self):
        self.client.force_login(self.user)

        first_page = self.client.get(reverse("schedule"))
        second_page = self.client.get(reverse("schedule"), {"page": 2})

        self.assertEqual(len(first_page.context["schedules"]), 20)
        self.assertEqual(len(second_page.context["schedules"]), 5)
        self.assertContains(first_page, "25 урок")
        self.assertContains(second_page, "5 урок")
        self.assertContains(first_page, 'aria-current="page"')
        self.assertContains(first_page, 'aria-label="Страница 2"')
        self.assertContains(second_page, 'aria-label="Страница 1"')

    def test_filter_is_applied_before_pagination(self):
        self.client.force_login(self.user)

        response = self.client.get(
            reverse("schedule"),
            {"lesson_type": "Новая тема"},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.context["schedules"]), 12)
        self.assertTrue(
            all(
                schedule.lesson_type == "Новая тема"
                for schedule in response.context["schedules"]
            )
        )

    def test_ajax_pagination_returns_only_schedule_fragment(self):
        self.client.force_login(self.user)

        response = self.client.get(
            reverse("schedule"),
            {"page": 2},
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        )

        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(
            response,
            "Course/includes/schedule_results.html",
        )
        self.assertContains(response, 'id="schedule-results"')
        self.assertContains(response, 'aria-current="page"')
        self.assertNotContains(response, "<html")
        self.assertEqual(len(response.context["schedules"]), 5)

    def test_timetable_responses_cannot_cache_ajax_fragment_as_full_page(self):
        self.client.force_login(self.user)
        schedule_url = reverse("schedule")

        full_response = self.client.get(schedule_url, {"page": 2})
        ajax_response = self.client.get(
            schedule_url,
            {"page": 2},
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        )

        self.assertContains(full_response, "<html")
        self.assertNotContains(ajax_response, "<html")
        for response in (full_response, ajax_response):
            self.assertIn("no-store", response.headers["Cache-Control"])
            self.assertIn("X-Requested-With", response.headers["Vary"])

    def test_ajax_filter_returns_only_matching_schedule_fragment(self):
        self.client.force_login(self.user)

        response = self.client.get(
            reverse("schedule"),
            {"lesson_type": "Новая тема"},
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        )

        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(
            response,
            "Course/includes/schedule_results.html",
        )
        self.assertNotContains(response, "<html")
        self.assertTrue(response.context["schedules"])
        self.assertTrue(
            all(
                schedule.lesson_type == "Новая тема"
                for schedule in response.context["schedules"]
            )
        )


class ScheduleOrderingTests(TimetableTests):
    def test_new_lesson_is_added_to_the_end_of_its_direction(self):
        lesson = Schedule.objects.create(
            theme='Новый урок',
            plan='План',
            lesson_materials='Материалы',
            direction=self.direction,
        )

        self.assertEqual(lesson.position, 26)

    def test_lesson_can_be_inserted_after_another_lesson(self):
        existing_lessons = list(
            Schedule.objects.filter(direction=self.direction).order_by('position', 'pk')
        )
        inserted_lesson = Schedule.objects.create(
            theme='Урок между темами',
            plan='План',
            lesson_materials='Материалы',
            direction=self.direction,
        )
        schedule_admin = ScheduleAdmin(Schedule, admin.site)

        with transaction.atomic():
            schedule_admin._place_after(inserted_lesson, existing_lessons[5])

        ordered_lessons = list(
            Schedule.objects.filter(direction=self.direction).order_by('position', 'pk')
        )
        self.assertEqual(ordered_lessons[6].pk, inserted_lesson.pk)
        self.assertEqual(
            [lesson.position for lesson in ordered_lessons],
            list(range(1, 27)),
        )

    def test_admin_add_form_inserts_lesson_after_selected_one(self):
        admin_user = User.objects.create_superuser(
            'schedule-admin',
            password='password',
        )
        after_lesson = list(
            Schedule.objects.filter(direction=self.direction).order_by('position', 'pk')
        )[3]
        self.client.force_login(admin_user)

        response = self.client.post(
            reverse('admin:Course_schedule_add'),
            {
                'theme': 'Урок из админки',
                'plan': 'План',
                'lesson_materials': 'Материалы',
                'lesson_type': 'Новая тема',
                'direction': self.direction.pk,
                'insert_after': after_lesson.pk,
                '_save': 'Сохранить',
            },
        )

        self.assertEqual(response.status_code, 302)
        inserted_lesson = Schedule.objects.get(theme='Урок из админки')
        ordered_ids = list(
            Schedule.objects.filter(direction=self.direction)
            .order_by('position', 'pk')
            .values_list('pk', flat=True)
        )
        self.assertEqual(
            ordered_ids[ordered_ids.index(after_lesson.pk) + 1],
            inserted_lesson.pk,
        )

    def test_archived_lesson_is_hidden_from_student_program(self):
        archived_lesson = Schedule.objects.filter(direction=self.direction).first()
        archived_lesson.is_archived = True
        archived_lesson.save(update_fields=['is_archived'])
        schedule_admin = ScheduleAdmin(Schedule, admin.site)

        with transaction.atomic():
            schedule_admin._normalize_direction(self.direction.pk)

        available_ids = get_accessible_schedule_ids(self.student)
        active_positions = list(
            Schedule.objects.filter(
                direction=self.direction,
                is_archived=False,
            ).values_list('position', flat=True)
        )
        self.assertNotIn(archived_lesson.pk, available_ids)
        self.assertEqual(active_positions, list(range(1, 25)))


class AdminFilterTests(TimetableTests):
    def test_active_filters_exclude_inactive_students_and_groups(self):
        inactive_group = LearnGroup.objects.create(
            title='Архивная группа',
            is_studies=False,
            teacher=self.student,
        )
        inactive_user = User.objects.create_user(
            'archived-student',
            email='archived@example.com',
            password='password',
        )
        inactive_student = Student.objects.create(
            user=inactive_user,
            contact='@archived',
            is_learned=False,
            groups=inactive_group,
        )
        request = RequestFactory().get('/admin/Course/lessonsolution/')
        request.user = User.objects.create_superuser(
            'filter-admin',
            password='password',
        )
        model_admin = LessonSolutionAdmin(LessonSolution, admin.site)

        student_filter = ActiveStudentListFilter(
            request,
            {},
            LessonSolution,
            model_admin,
        )
        group_filter = ActiveGroupListFilter(
            request,
            {},
            LessonSolution,
            model_admin,
        )

        self.assertIn(
            (self.student.pk, self.student.user.username),
            student_filter.lookups(request, model_admin),
        )
        self.assertNotIn(
            (inactive_student.pk, inactive_student.user.username),
            student_filter.lookups(request, model_admin),
        )
        self.assertIn(
            (self.group.pk, self.group.title),
            group_filter.lookups(request, model_admin),
        )
        self.assertNotIn(
            (inactive_group.pk, inactive_group.title),
            group_filter.lookups(request, model_admin),
        )

    def test_solution_quick_filters_and_my_groups(self):
        schedules = list(Schedule.objects.filter(direction=self.direction)[:3])
        pending_solution = LessonSolution.objects.create(
            student=self.student,
            schedule=schedules[0],
            status=LessonSolution.Status.PENDING,
        )
        revision_solution = LessonSolution.objects.create(
            student=self.student,
            schedule=schedules[1],
            status=LessonSolution.Status.NEEDS_REVISION,
        )
        accepted_solution = LessonSolution.objects.create(
            student=self.student,
            schedule=schedules[2],
            status=LessonSolution.Status.ACCEPTED,
            teacher_comment='Всё хорошо.',
        )
        LessonSolution.objects.filter(pk=accepted_solution.pk).update(
            submitted_at=timezone.now() - timedelta(days=8),
        )
        request = RequestFactory().get('/admin/Course/lessonsolution/')
        request.user = self.user
        model_admin = LessonSolutionAdmin(LessonSolution, admin.site)

        def apply_filter(filter_class, parameter_name, value):
            filter_instance = filter_class(
                request,
                {parameter_name: [value]},
                LessonSolution,
                model_admin,
            )
            return set(
                filter_instance.queryset(
                    request,
                    LessonSolution.objects.all(),
                ).values_list('pk', flat=True),
            )

        self.assertEqual(
            apply_filter(AttentionNeededListFilter, 'requires_attention', 'yes'),
            {pending_solution.pk, revision_solution.pk},
        )
        self.assertEqual(
            apply_filter(WithoutCommentListFilter, 'without_comment', 'yes'),
            {pending_solution.pk, revision_solution.pk},
        )
        self.assertEqual(
            apply_filter(RecentSolutionListFilter, 'recent_submission', 'week'),
            {pending_solution.pk, revision_solution.pk},
        )
        self.assertEqual(
            apply_filter(MyGroupsListFilter, 'my_groups', 'yes'),
            {pending_solution.pk, revision_solution.pk, accepted_solution.pk},
        )


class LessonSolutionTests(TimetableTests):
    def setUp(self):
        super().setUp()
        self.media_directory = tempfile.TemporaryDirectory()
        self.settings_override = override_settings(
            PRIVATE_SOLUTION_MEDIA_ROOT=self.media_directory.name,
        )
        self.settings_override.enable()
        self.addCleanup(self.settings_override.disable)
        self.addCleanup(self.media_directory.cleanup)
        self.schedule = Schedule.objects.order_by('pk').first()
        self.upload_url = reverse(
            'lesson_solution_upload',
            args=(self.schedule.pk,),
        )

    def _upload(self, *files):
        return self.client.post(
            self.upload_url,
            {
                'files': list(files),
                'next': reverse('schedule'),
            },
        )

    def test_student_can_upload_several_files_for_an_available_lesson(self):
        self.client.force_login(self.user)

        response = self._upload(
            SimpleUploadedFile('solution.py', b'print("done")'),
            SimpleUploadedFile('notes.md', b'# Explanation'),
        )

        self.assertRedirects(response, reverse('schedule'))
        solution = LessonSolution.objects.get(
            student=self.student,
            schedule=self.schedule,
        )
        self.assertEqual(solution.status, LessonSolution.Status.PENDING)
        self.assertEqual(solution.files.count(), 2)
        submission = solution.submissions.get()
        self.assertEqual(submission.attempt_number, 1)
        self.assertTrue(
            solution.files.first().file.name.startswith(
                f'solution_uploads/{self.student.pk}/',
            )
        )
        self.assertNotIn('solution.py', solution.files.first().file.name)

    def test_resubmission_replaces_files_and_resets_review_status(self):
        self.client.force_login(self.user)
        self._upload(SimpleUploadedFile('first.py', b'print(1)'))
        solution = LessonSolution.objects.get(
            student=self.student,
            schedule=self.schedule,
        )
        old_file = solution.files.get()
        old_file_name = old_file.file.name
        old_file_storage = old_file.file.storage
        reviewer = User.objects.create_superuser(
            'reviewer',
            password='password',
        )
        solution.status = LessonSolution.Status.NEEDS_REVISION
        solution.teacher_comment = 'Добавь обработку ошибки.'
        solution.reviewed_by = reviewer
        solution.save()

        response = self._upload(SimpleUploadedFile('second.py', b'print(2)'))

        self.assertRedirects(response, reverse('schedule'))
        solution.refresh_from_db()
        self.assertEqual(solution.status, LessonSolution.Status.PENDING)
        self.assertEqual(solution.teacher_comment, '')
        self.assertIsNone(solution.reviewed_by)
        self.assertEqual(solution.files.count(), 1)
        self.assertEqual(solution.files.get().original_name, 'second.py')
        self.assertEqual(
            list(
                solution.submissions.order_by('attempt_number').values_list(
                    'attempt_number',
                    flat=True,
                )
            ),
            [1, 2],
        )
        self.assertFalse(old_file_storage.exists(old_file_name))

    def test_upload_rejects_an_unavailable_lesson_and_unsupported_file(self):
        self.client.force_login(self.user)
        other_direction = DirectionStudy.objects.create(title='Java')
        unavailable_schedule = Schedule.objects.create(
            theme='Недоступный урок',
            plan='План',
            lesson_materials='Материалы',
            direction=other_direction,
        )

        unavailable_response = self.client.post(
            reverse('lesson_solution_upload', args=(unavailable_schedule.pk,)),
            {'files': [SimpleUploadedFile('solution.py', b'print(1)')]},
        )
        invalid_file_response = self._upload(
            SimpleUploadedFile('program.exe', b'not executable'),
        )

        self.assertEqual(unavailable_response.status_code, 403)
        self.assertRedirects(invalid_file_response, reverse('schedule'))
        self.assertFalse(
            LessonSolution.objects.filter(
                student=self.student,
                schedule=self.schedule,
            ).exists()
        )

    def test_solution_file_is_available_only_to_owner_or_superuser(self):
        self.client.force_login(self.user)
        self._upload(SimpleUploadedFile('solution.py', b'print("private")'))
        solution_file = LessonSolution.objects.get(
            student=self.student,
            schedule=self.schedule,
        ).files.get()
        download_url = reverse(
            'lesson_solution_file_download',
            args=(solution_file.pk,),
        )

        owner_response = self.client.get(download_url)

        self.assertEqual(owner_response.status_code, 200)
        self.assertIn('attachment;', owner_response['Content-Disposition'])
        self.assertEqual(
            b''.join(owner_response.streaming_content),
            b'print("private")',
        )

        preview_response = self.client.get(f'{download_url}?view=1')

        self.assertEqual(preview_response.status_code, 200)
        self.assertNotIn('attachment;', preview_response.get('Content-Disposition', ''))
        self.assertEqual(
            preview_response['Content-Type'],
            'text/plain; charset=utf-8',
        )
        self.assertEqual(preview_response['X-Content-Type-Options'], 'nosniff')

        other_user = User.objects.create_user(
            'other-student', email='other@example.com', password='password',
        )
        other_student = Student.objects.create(
            user=other_user,
            contact='@other',
            is_learned=True,
            groups=self.group,
        )
        other_student.direction.add(self.direction)
        self.client.force_login(other_user)
        other_response = self.client.get(download_url)

        self.assertEqual(other_response.status_code, 403)

        superuser = User.objects.create_superuser('admin', password='password')
        self.client.force_login(superuser)
        admin_response = self.client.get(download_url)

        self.assertEqual(admin_response.status_code, 200)

    def test_solution_admin_page_shows_attached_file_links(self):
        self.client.force_login(self.user)
        self._upload(SimpleUploadedFile('solution.py', b'print("shown")'))
        solution = LessonSolution.objects.get(
            student=self.student,
            schedule=self.schedule,
        )
        admin_user = User.objects.create_superuser('files-admin', password='password')
        self.client.force_login(admin_user)

        response = self.client.get(
            reverse('admin:Course_lessonsolution_change', args=(solution.pk,)),
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'solution.py')
        self.assertContains(response, 'Открыть')
        self.assertContains(response, 'Скачать')
        self.assertContains(response, 'История отправок решений')
        self.assertContains(
            response,
            reverse('py_interpreter_solution_file', args=(solution.files.get().pk,)),
        )

    def test_teacher_opens_python_solution_in_interpreter(self):
        self.client.force_login(self.user)
        code = 'print("Привет")\n# </textarea><script>alert(1)</script>'
        self._upload(SimpleUploadedFile('solution.py', code.encode('utf-8')))
        solution_file = LessonSolution.objects.get(student=self.student).files.get()
        url = reverse('py_interpreter_solution_file', args=(solution_file.pk,))
        self.user.is_staff = True
        self.user.save(update_fields=['is_staff'])

        response = self.client.get(url)

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['source_code'], code)
        self.assertContains(response, 'data-solution-file-id=')
        self.assertContains(response, 'solution.py')
        self.assertContains(response, '&lt;/textarea&gt;&lt;script&gt;')
        self.assertNotContains(response, '<script>alert(1)</script>')
        self.assertIn('no-store', response['Cache-Control'])

    def test_student_and_unrelated_teacher_cannot_open_solution_in_interpreter(self):
        self.client.force_login(self.user)
        self._upload(SimpleUploadedFile('solution.py', b'print(1)'))
        solution_file = LessonSolution.objects.get(student=self.student).files.get()
        url = reverse('py_interpreter_solution_file', args=(solution_file.pk,))

        self.assertEqual(self.client.get(url).status_code, 403)

        other_teacher = User.objects.create_user('other-teacher', password='password', is_staff=True)
        self.client.force_login(other_teacher)
        self.assertEqual(self.client.get(url).status_code, 403)

        self.client.logout()
        self.assertEqual(self.client.get(url).status_code, 302)

        superuser = User.objects.create_superuser('solutions-admin', password='password')
        self.client.force_login(superuser)
        self.assertEqual(self.client.get(url).status_code, 200)

    def test_notebook_code_is_loaded_without_metadata_or_markdown(self):
        self.client.force_login(self.user)
        notebook = b'{"cells": [{"cell_type": "markdown", "source": ["private notes"]}, {"cell_type": "code", "source": ["print(1)\\n", "print(2)"]}], "metadata": {"secret": "not in editor"}}'
        self._upload(SimpleUploadedFile('work.ipynb', notebook))
        solution_file = LessonSolution.objects.get(student=self.student).files.get()
        superuser = User.objects.create_superuser('notebook-admin', password='password')
        self.client.force_login(superuser)

        response = self.client.get(reverse('py_interpreter_solution_file', args=(solution_file.pk,)))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['source_code'], 'print(1)\nprint(2)')
        self.assertEqual(response.context['download_name'], 'work.py')
        self.assertNotContains(response, 'private notes')

    def test_non_code_file_offers_original_without_rendering_binary(self):
        self.client.force_login(self.user)
        self._upload(SimpleUploadedFile('work.pdf', b'%PDF-1.7\x00test'))
        solution_file = LessonSolution.objects.get(student=self.student).files.get()
        superuser = User.objects.create_superuser('pdf-admin', password='password')
        self.client.force_login(superuser)

        response = self.client.get(reverse('py_interpreter_solution_file', args=(solution_file.pk,)))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Этот формат нельзя открыть как Python-код')
        self.assertContains(response, 'Скачать оригинал')
        self.assertNotContains(response, '%PDF-1.7')

    def test_long_source_is_not_silently_truncated_in_editor(self):
        self.client.force_login(self.user)
        self._upload(SimpleUploadedFile('long.py', b'#' * 20001))
        solution_file = LessonSolution.objects.get(student=self.student).files.get()
        superuser = User.objects.create_superuser('long-source-admin', password='password')
        self.client.force_login(superuser)

        response = self.client.get(reverse('py_interpreter_solution_file', args=(solution_file.pk,)))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context['source_code'], '')
        self.assertContains(response, 'Код превышает лимит редактора')
        self.assertContains(response, 'Скачать оригинал')

    @override_settings(PROXYAPI_API_KEY='test-key', PROXYAPI_LESSON_SOLUTION_MODEL='gpt-4o-mini')
    def test_upload_creates_ai_suggestion_and_admin_can_send_it_for_revision(self):
        self.client.force_login(self.user)
        secret = 'sk_abcdefghijklmnopqrstuvwxyz123456'
        response_payload = {
            'choices': [{'message': {'content': json.dumps({
                'summary': 'Есть ошибка в обработке пустого списка.',
                'issues': [{'file': 'solution.py', 'problem': 'solution.py: Учтите пустой список.'}],
            })}}],
        }
        with patch('Course.lesson_solution_ai.requests.post') as ai_post:
            ai_post.return_value.json.return_value = response_payload
            with self.captureOnCommitCallbacks(execute=True):
                upload_response = self._upload(SimpleUploadedFile(
                    'solution.py', f'API_KEY = "{secret}"\nprint(1)'.encode(),
                ))

        self.assertEqual(upload_response.status_code, 302)
        submission = LessonSolutionSubmission.objects.get()
        self.assertEqual(submission.ai_status, submission.AIStatus.READY)
        self.assertEqual(submission.ai_issue_count, 1)
        self.assertEqual(submission.ai_feedback, '1. Учтите пустой список.')
        request_payload = ai_post.call_args.kwargs['json']
        self.assertEqual(request_payload['model'], 'gpt-4o-mini')
        review_rules = request_payload['messages'][0]['content']
        self.assertIn('решение начинающего ученика', review_rules)
        self.assertIn('обычном допустимом примере', review_rules)
        self.assertIn('Не считай ошибками отсутствие комментариев', review_rules)
        self.assertIn('Если существенных ошибок нет, верни пустой issues', review_rules)
        self.assertIn('<REDACTED>', request_payload['messages'][1]['content'])
        self.assertNotIn(secret, request_payload['messages'][1]['content'])

        student_page = self.client.get(reverse('schedule'))
        self.assertNotContains(student_page, 'Учтите пустой список.')

        admin_user = User.objects.create_superuser('ai-admin', password='password')
        self.client.force_login(admin_user)
        admin_url = reverse('admin:Course_lessonsolution_change', args=(submission.solution_id,))
        admin_page = self.client.get(admin_url)
        self.assertContains(admin_page, 'Учтите пустой список.')
        self.assertContains(admin_page, 'Отправить на доработку с этим текстом')

        apply_response = self.client.post(
            reverse('admin:Course_lessonsolution_apply_ai_feedback', args=(submission.solution_id,)),
            {'submission_id': submission.pk},
        )
        self.assertRedirects(apply_response, admin_url)
        solution = submission.solution
        solution.refresh_from_db()
        self.assertEqual(solution.status, LessonSolution.Status.NEEDS_REVISION)
        self.assertEqual(solution.teacher_comment, submission.ai_feedback)
        self.assertEqual(solution.reviewed_by, admin_user)

    @override_settings(PROXYAPI_API_KEY='test-key')
    def test_ai_without_issues_does_not_offer_revision_button(self):
        self.client.force_login(self.user)
        with patch('Course.lesson_solution_ai.requests.post') as ai_post:
            ai_post.return_value.json.return_value = {
                'choices': [{'message': {'content': '{"summary":"OK","issues":[]}'}}],
            }
            with self.captureOnCommitCallbacks(execute=True):
                self._upload(SimpleUploadedFile('solution.py', b'print(1)'))

        submission = LessonSolutionSubmission.objects.get()
        self.assertEqual(submission.ai_status, submission.AIStatus.READY)
        self.assertEqual(submission.ai_issue_count, 0)
        self.assertIn('явных ошибок не найдено', submission.ai_feedback)
        admin_user = User.objects.create_superuser('no-issue-admin', password='password')
        self.client.force_login(admin_user)
        page = self.client.get(reverse(
            'admin:Course_lessonsolution_change', args=(submission.solution_id,),
        ))
        self.assertNotContains(page, 'Отправить на доработку с этим текстом')
        response = self.client.post(
            reverse('admin:Course_lessonsolution_apply_ai_feedback', args=(submission.solution_id,)),
            {'submission_id': submission.pk},
        )
        self.assertEqual(response.status_code, 302)
        submission.solution.refresh_from_db()
        self.assertEqual(submission.solution.status, LessonSolution.Status.PENDING)

    @override_settings(PROXYAPI_API_KEY='test-key')
    def test_unsupported_file_is_not_reported_as_error_free(self):
        self.client.force_login(self.user)
        with patch('Course.lesson_solution_ai.requests.post') as ai_post:
            with self.captureOnCommitCallbacks(execute=True):
                self._upload(SimpleUploadedFile('answer.pdf', b'%PDF-1.7\x00data'))

        submission = LessonSolutionSubmission.objects.get()
        self.assertEqual(submission.ai_status, submission.AIStatus.UNSUPPORTED)
        self.assertIn('answer.pdf', submission.ai_source_summary)
        self.assertEqual(submission.ai_feedback, '')
        ai_post.assert_not_called()

    @override_settings(PROXYAPI_API_KEY='test-key')
    def test_mixed_files_are_marked_partially_reviewed(self):
        self.client.force_login(self.user)
        with patch('Course.lesson_solution_ai.requests.post') as ai_post:
            ai_post.return_value.json.return_value = {
                'choices': [{'message': {'content': '{"summary":"OK","issues":[]}'}}],
            }
            with self.captureOnCommitCallbacks(execute=True):
                self._upload(
                    SimpleUploadedFile('solution.py', b'print(1)'),
                    SimpleUploadedFile('diagram.pdf', b'%PDF-1.7\x00data'),
                )

        submission = LessonSolutionSubmission.objects.get()
        self.assertEqual(submission.ai_status, submission.AIStatus.PARTIAL)
        self.assertIn('solution.py', submission.ai_source_summary)
        self.assertIn('diagram.pdf', submission.ai_source_summary)
        self.assertIn('Полную проверку провести не удалось', submission.ai_feedback)
        self.assertNotIn('%PDF', ai_post.call_args.kwargs['json']['messages'][1]['content'])

    @override_settings(PROXYAPI_API_KEY='test-key')
    def test_docx_text_is_included_in_ai_review(self):
        from docx import Document

        document = Document()
        document.add_paragraph('Ответ находится в документе.')
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = 'Проверить граничный случай.'
        buffer = BytesIO()
        document.save(buffer)
        self.client.force_login(self.user)
        with patch('Course.lesson_solution_ai.requests.post') as ai_post:
            ai_post.return_value.json.return_value = {
                'choices': [{'message': {'content': '{"summary":"OK","issues":[]}'}}],
            }
            with self.captureOnCommitCallbacks(execute=True):
                self._upload(SimpleUploadedFile('answer.docx', buffer.getvalue()))

        submission = LessonSolutionSubmission.objects.get()
        self.assertEqual(submission.ai_status, submission.AIStatus.READY)
        ai_input = ai_post.call_args.kwargs['json']['messages'][1]['content']
        self.assertIn('Ответ находится в документе.', ai_input)
        self.assertIn('Проверить граничный случай.', ai_input)

    @override_settings(PROXYAPI_API_KEY='test-key')
    def test_old_ai_recommendation_cannot_be_sent_after_new_submission(self):
        self.client.force_login(self.user)
        with patch('Course.lesson_solution_ai.requests.post') as ai_post:
            ai_post.return_value.json.return_value = {
                'choices': [{'message': {'content': json.dumps({
                    'summary': 'Ошибка',
                    'issues': [{'file': 'first.py', 'problem': 'Исправьте условие.'}],
                })}}],
            }
            with self.captureOnCommitCallbacks(execute=True):
                self._upload(SimpleUploadedFile('first.py', b'print(1)'))
        first_submission = LessonSolutionSubmission.objects.get()
        self._upload(SimpleUploadedFile('second.py', b'print(2)'))

        admin_user = User.objects.create_superuser('stale-admin', password='password')
        self.client.force_login(admin_user)
        admin_url = reverse('admin:Course_lessonsolution_change', args=(first_submission.solution_id,))
        page = self.client.get(admin_url)
        self.assertNotContains(page, 'Исправьте условие.')
        response = self.client.post(
            reverse('admin:Course_lessonsolution_apply_ai_feedback', args=(first_submission.solution_id,)),
            {'submission_id': first_submission.pk},
        )
        self.assertRedirects(response, admin_url)
        first_submission.solution.refresh_from_db()
        self.assertEqual(first_submission.solution.status, LessonSolution.Status.PENDING)
        self.assertEqual(first_submission.solution.teacher_comment, '')

    @override_settings(PROXYAPI_API_KEY='test-key')
    def test_ai_failure_does_not_block_upload_and_can_be_retried(self):
        self.client.force_login(self.user)
        with patch('Course.lesson_solution_ai.requests.post', side_effect=requests.Timeout):
            with self.captureOnCommitCallbacks(execute=True):
                upload_response = self._upload(SimpleUploadedFile('solution.py', b'print(1)'))

        self.assertEqual(upload_response.status_code, 302)
        submission = LessonSolutionSubmission.objects.get()
        self.assertEqual(submission.ai_status, submission.AIStatus.FAILED)
        admin_user = User.objects.create_superuser('retry-admin', password='password')
        self.client.force_login(admin_user)
        admin_page = self.client.get(reverse(
            'admin:Course_lessonsolution_change', args=(submission.solution_id,),
        ))
        self.assertContains(admin_page, 'Повторить проверку ИИ')
        with patch('Course.lesson_solution_ai.requests.post') as ai_post:
            ai_post.return_value.json.return_value = {
                'choices': [{'message': {'content': '{"summary":"OK","issues":[]}'}}],
            }
            retry_response = self.client.post(
                reverse('admin:Course_lessonsolution_retry_ai_review', args=(submission.solution_id,)),
                {'submission_id': submission.pk},
            )

        self.assertEqual(retry_response.status_code, 302)
        submission.refresh_from_db()
        self.assertEqual(submission.ai_status, submission.AIStatus.READY)

    def test_other_teacher_cannot_apply_ai_feedback(self):
        self.client.force_login(self.user)
        self._upload(SimpleUploadedFile('solution.py', b'print(1)'))
        submission = LessonSolutionSubmission.objects.get()
        submission.ai_status = submission.AIStatus.READY
        submission.ai_feedback = 'Нужно исправить решение.'
        submission.ai_issue_count = 1
        submission.save(update_fields=('ai_status', 'ai_feedback', 'ai_issue_count'))
        other_teacher = User.objects.create_user('unrelated-teacher', is_staff=True)
        other_teacher.user_permissions.add(Permission.objects.get(
            content_type__app_label='Course',
            codename='change_lessonsolution',
        ))
        self.client.force_login(other_teacher)

        action_url = reverse(
            'admin:Course_lessonsolution_apply_ai_feedback',
            args=(submission.solution_id,),
        )
        self.assertEqual(self.client.get(action_url).status_code, 405)
        self.assertEqual(
            self.client.post(action_url, {'submission_id': submission.pk}).status_code,
            404,
        )
        submission.solution.refresh_from_db()
        self.assertEqual(submission.solution.status, LessonSolution.Status.PENDING)

    @override_settings(PROXYAPI_API_KEY='test-key')
    def test_stuck_pending_ai_review_can_be_retried(self):
        self.client.force_login(self.user)
        self._upload(SimpleUploadedFile('solution.py', b'print(1)'))
        submission = LessonSolutionSubmission.objects.get()
        LessonSolutionSubmission.objects.filter(pk=submission.pk).update(
            submitted_at=timezone.now() - timedelta(minutes=3),
        )
        admin_user = User.objects.create_superuser('stuck-admin', password='password')
        self.client.force_login(admin_user)
        admin_url = reverse('admin:Course_lessonsolution_change', args=(submission.solution_id,))
        self.assertContains(self.client.get(admin_url), 'Повторить проверку ИИ')
        with patch('Course.lesson_solution_ai.requests.post') as ai_post:
            ai_post.return_value.json.return_value = {
                'choices': [{'message': {'content': '{"summary":"OK","issues":[]}'}}],
            }
            response = self.client.post(
                reverse('admin:Course_lessonsolution_retry_ai_review', args=(submission.solution_id,)),
                {'submission_id': submission.pk},
            )
        self.assertRedirects(response, admin_url)
        submission.refresh_from_db()
        self.assertEqual(submission.ai_status, submission.AIStatus.READY)

    def test_old_ai_feedback_loses_file_names_without_overwriting_manual_edits(self):
        self.client.force_login(self.user)
        self._upload(SimpleUploadedFile('a.py', b'print(1)'))
        submission = LessonSolutionSubmission.objects.get()
        old_feedback = (
            'По предварительной проверке есть замечания:\n'
            '1. a.py: Исправьте условие.\n'
            '2. a.py: Добавьте обработку пустого списка.'
        )
        submission.ai_feedback = old_feedback
        submission.save(update_fields=('ai_feedback',))
        solution = submission.solution
        solution.teacher_comment = old_feedback.replace('\n', '\r\n')
        solution.save(update_fields=('teacher_comment',))

        another_solution = LessonSolution.objects.create(
            student=self.student,
            schedule=Schedule.objects.exclude(pk=self.schedule.pk).first(),
            teacher_comment=old_feedback + '\nМой комментарий.',
        )
        another_submission = LessonSolutionSubmission.objects.create(
            solution=another_solution,
            attempt_number=1,
            submitted_at=timezone.now(),
            ai_feedback=old_feedback,
        )
        migration = importlib.import_module(
            'Course.migrations.0027_clean_lesson_solution_ai_feedback',
        )
        migration.clean_existing_ai_feedback(apps, None)

        submission.refresh_from_db()
        solution.refresh_from_db()
        another_submission.refresh_from_db()
        another_solution.refresh_from_db()
        expected = (
            '1. Исправьте условие.\n'
            '2. Добавьте обработку пустого списка.'
        )
        self.assertEqual(submission.ai_feedback, expected)
        self.assertEqual(solution.teacher_comment, expected.replace('\n', '\r\n'))
        self.assertEqual(another_submission.ai_feedback, expected)
        self.assertEqual(
            another_solution.teacher_comment,
            old_feedback + '\nМой комментарий.',
        )

    def test_solution_list_prioritizes_pending_reviews(self):
        self.client.force_login(self.user)
        self._upload(SimpleUploadedFile('pending.py', b'print("pending")'))
        pending_solution = LessonSolution.objects.get(
            student=self.student,
            schedule=self.schedule,
        )
        accepted_schedule = Schedule.objects.exclude(pk=self.schedule.pk).first()
        accepted_solution = LessonSolution.objects.create(
            student=self.student,
            schedule=accepted_schedule,
            status=LessonSolution.Status.ACCEPTED,
        )
        admin_user = User.objects.create_superuser('queue-admin', password='password')
        self.client.force_login(admin_user)

        response = self.client.get(reverse('admin:Course_lessonsolution_changelist'))
        page = response.content.decode()

        self.assertEqual(response.status_code, 200)
        self.assertLess(
            page.find(pending_solution.schedule.theme),
            page.find(accepted_solution.schedule.theme),
        )
        self.assertNotContains(response, 'Проверено')
        self.assertNotContains(response, 'Файлов')
        change_url = reverse(
            'admin:Course_lessonsolution_change',
            args=(pending_solution.pk,),
        )
        self.assertContains(response, f'href="{change_url}"', count=5)


class AccountLoginTests(TimetableTests):
    def _login(self, login):
        response = self.client.post(
            reverse('account_login'),
            {'login': login, 'password': 'password'},
        )

        self.assertEqual(response.status_code, 302)
        self.assertEqual(int(self.client.session['_auth_user_id']), self.user.pk)

    def test_user_can_log_in_with_username(self):
        self._login(self.user.username)

    def test_student_can_log_in_with_account_email(self):
        self._login(self.user.email)
